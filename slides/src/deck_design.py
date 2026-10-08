"""Deck design system for HSI-Learning lecture slides (python-pptx).

Editorial, restrained layout: ink-on-paper typography, one terracotta
accent, a hyperspectral gradient strip as the course identity motif,
real typeset formulas (matplotlib mathtext -> transparent PNG), and
hairline "term spines" instead of decorative card stacks.

Every run sets BOTH the latin and East-Asian typeface so Chinese text
never falls back to a theme default. No footers, no page numbers.
"""
from __future__ import annotations

import hashlib
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
matplotlib.rcParams['mathtext.fontset'] = 'cm'
import matplotlib.pyplot as plt
from PIL import Image

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- tokens
INK = '1C2B3A'        # titles + body
SUB = '5C6B7A'        # secondary text
LINE = 'D9E2E9'       # hairlines
PANEL = 'F5F8FA'      # quiet panel fill
ACCENT = 'C9571E'     # terracotta accent
CODE_BG = '202E3C'
CODE_FG = 'DEE9F2'
CODE_COMMENT = '7E99AD'
CJK = 'Microsoft YaHei'
MONO = 'Consolas'

# Visible-spectrum inspired stops, violet -> red (course identity motif).
SPECTRAL_STOPS = ['472A8A', '2C5FB0', '1F8FB8', '2E9E4F', 'F2C316', 'E8752A', 'D33A2C']

SLIDE_W, SLIDE_H = 13.333, 7.5


# ---------------------------------------------------------------- low level
def _rgb(hexstr):
    return RGBColor.from_string(hexstr)


def _no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def _style_run(run, size, color=INK, bold=False, font=CJK, tracking=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = font
    f.color.rgb = _rgb(color)
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn('a:latin'))
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = parse_xml(f'<a:ea {nsdecls("a")} typeface="{CJK}"/>')
        if latin is not None:
            latin.addnext(ea)
        else:
            rPr.append(ea)
    if tracking:
        rPr.set('spc', str(tracking))


def _tf(slide, *args):
    """Add or fetch a textbox; returns (shape, text_frame)."""
    if len(args) == 1:  # existing shape
        shape = args[0]
        return shape, shape.text_frame
    x, y, w, h = args
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    return box, tf


def _para(tf, runs, *, first=False, align=None, space_after=None, space_before=None,
          line=None, size=18, color=INK, bold=False, font=CJK):
    """runs: str, or list of (text, overrides-dict)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.space_after = Pt(space_after)
    if space_before is not None:
        p.space_before = Pt(space_before)
    if line is not None:
        p.line_spacing = line
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, over in runs:
        run = p.add_run()
        run.text = text
        _style_run(run, over.get('size', size), over.get('color', color),
                   over.get('bold', bold), over.get('font', font),
                   over.get('tracking'))
    return p


def text(slide, x, y, w, h, lines, **kw):
    """lines: list of run-specs (str or list of tuples); one paragraph each."""
    box, tf = _tf(slide, x, y, w, h)
    for i, spec in enumerate(lines):
        _para(tf, spec, first=(i == 0), **kw)
    return box


def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=0.75,
         shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _no_shadow(sp)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(fill)
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = _rgb(line_color)
        sp.line.width = Pt(line_w)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    sp.text_frame.paragraphs[0].text = ''
    return sp


def gradient_bar(slide, x, y, w, h, stops=SPECTRAL_STOPS, angle_deg=90):
    sp = rect(slide, x, y, w, h)
    spPr = sp._element.spPr
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill'):
        el = spPr.find(qn(tag))
        if el is not None:
            spPr.remove(el)
    gs = ''.join(
        f'<a:gs pos="{int(round(i * 100000 / (len(stops) - 1)))}"><a:srgbClr val="{c}"/></a:gs>'
        for i, c in enumerate(stops))
    frag = parse_xml(
        f'<a:gradFill {nsdecls("a")} rotWithShape="1"><a:gsLst>{gs}</a:gsLst>'
        f'<a:lin ang="{int(angle_deg * 60000)}" scaled="1"/></a:gradFill>')
    geom = spPr.find(qn('a:prstGeom'))
    if geom is not None:
        geom.addnext(frag)
    else:
        spPr.append(frag)
    return sp


def dot(slide, cx, cy, d=0.11, color=ACCENT):
    return rect(slide, cx - d / 2, cy - d / 2, d, d, fill=color, shape=MSO_SHAPE.OVAL)


def hairline(slide, x, y, w, color=LINE, h=0.013):
    return rect(slide, x, y, w, h, fill=color)


# ---------------------------------------------------------------- formulas
_FORMULA_DPI = 300
_FORMULA_PT = 30


def formula_png(latex, cache_dir: Path, color='#1C2B3A'):
    """Render mathtext to a transparent high-DPI PNG; returns (path, w_in, h_in)."""
    cache_dir.mkdir(parents=True, exist_ok=True)
    key = hashlib.sha1(f'v2|{_FORMULA_PT}|{_FORMULA_DPI}|{latex}|{color}'.encode()).hexdigest()[:12]
    path = cache_dir / f'formula-{key}.png'
    if not path.exists():
        fig = plt.figure(figsize=(0.1, 0.1))
        fig.patch.set_alpha(0)
        fig.text(0, 0, f'${latex}$', fontsize=_FORMULA_PT, color=color)
        fig.savefig(path, dpi=_FORMULA_DPI, transparent=True,
                    bbox_inches='tight', pad_inches=0.03)
        plt.close(fig)
    with Image.open(path) as im:
        w_in, h_in = im.size[0] / _FORMULA_DPI, im.size[1] / _FORMULA_DPI
    return path, w_in, h_in


def place_formula(slide, latex, cache_dir: Path, cx, cy, max_w, max_h, color='#1C2B3A'):
    path, w, h = formula_png(latex, cache_dir, color)
    scale = min(max_w / w, max_h / h, 1.0)
    w, h = w * scale, h * scale
    slide.shapes.add_picture(str(path), Inches(cx - w / 2), Inches(cy - h / 2),
                             width=Inches(w), height=Inches(h))


# ---------------------------------------------------------------- blocks
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def header(slide, lesson, title):
    gradient_bar(slide, 0, 0, 0.09, SLIDE_H)
    text(slide, 0.92, 0.52, 11.5, 0.3,
         [[(lesson['id'], dict(bold=True, color=ACCENT, tracking=140)),
           ('  ·  ' + lesson['subtitle'], dict(color=SUB, tracking=140))]],
         size=11.5)
    text(slide, 0.92, 0.88, 11.5, 0.75, [title], size=28, bold=True)
    hairline(slide, 0.92, 1.8, 11.49)


def term_spine(slide, labels, x, y0, y1, label_size=16.5):
    """Vertical hairline with accent dots and terms — replaces card stacks."""
    labels = labels[:4]
    n = len(labels)
    ys = [y0 + i * (y1 - y0) / max(n - 1, 1) if n > 1 else (y0 + y1) / 2
          for i in range(n)]
    rect(slide, x, ys[0], 0.016, ys[-1] - ys[0] + 0.02, fill=LINE)
    for label, cy in zip(labels, ys):
        dot(slide, x + 0.008, cy)
        text(slide, x + 0.32, cy - 0.21, 4.9, 0.44, [label], size=label_size)


# ---------------------------------------------------------------- slides
def cover(prs, lesson):
    slide = blank(prs)
    gradient_bar(slide, 0, 0, 0.17, SLIDE_H)
    text(slide, 1.1, 1.12, 10.5, 0.35,
         [[('HSI-LEARNING', dict(tracking=300, bold=True)),
           ('  ·  高光谱图像分类课程', dict(tracking=200))]], size=12.5, color=SUB)
    text(slide, 1.08, 1.72, 3, 1, [lesson['id']], size=46, bold=True, color=ACCENT)
    text(slide, 1.08, 2.86, 10.9, 1.5, [lesson['title']], size=37, bold=True)
    text(slide, 1.1, 4.42, 10.5, 0.5, [lesson['subtitle']], size=18, color=SUB)
    hairline(slide, 1.1, 5.18, 10.9)
    text(slide, 1.1, 5.46, 10.9, 0.4,
         [[('本课交付 · ', dict(bold=True, color=INK)),
           ('  ·  '.join(lesson['outputs'][:3]), dict(color=SUB))]], size=15)
    text(slide, 1.1, 6.42, 10.9, 0.4,
         [[('先修 · ', dict(color=SUB)), (lesson['prerequisites'], dict(color=SUB))]],
         size=13)
    return slide


def objectives(prs, lesson):
    slide = blank(prs)
    header(slide, lesson, '本课交付')
    ys = [2.42, 3.62, 4.82]
    for i, (out, cy) in enumerate(zip(lesson['outputs'][:3], ys)):
        text(slide, 0.92, cy - 0.26, 0.85, 0.5, [f'0{i + 1}'], size=22,
             bold=True, color=ACCENT)
        text(slide, 1.82, cy - 0.22, 5.6, 0.6, [out], size=20)
        if i < 2:
            hairline(slide, 0.92, cy + 0.62, 6.5)
    rect(slide, 8.15, 2.2, 4.25, 3.5, fill=PANEL)
    text(slide, 8.5, 2.5, 3.6, 0.3, ['课堂入口'], size=11.5, color=ACCENT,
         bold=True)
    entry = [('NOTEBOOK', lesson['notebook'] or 'run_config.json / 精读笔记', MONO),
             ('CELLS', ', '.join(lesson['cells']), MONO),
             ('任务', f"{lesson['assignment']} · 交付：代码 + 配置 + 解释", None)]
    y = 3.02
    for label, value, font in entry:
        text(slide, 8.5, y, 3.55, 0.26, [label], size=10, color=SUB)
        text(slide, 8.5, y + 0.28, 3.55, 0.62,
             [[(value, dict(font=font, size=11) if font else {})]],
             size=12.5, line=1.25)
        y += 0.92
    return slide


def theory(prs, lesson, page, cache_dir):
    slide = blank(prs)
    header(slide, lesson, page['title'])
    lines = [[('—  ', dict(color=ACCENT, bold=True)), (pt, {})] for pt in page['points']]
    box = text(slide, 0.92, 2.25, 5.45, 3.3, lines, size=19, line=1.2,
               space_after=13)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if page.get('run'):
        code, out = page['run']
        run_size = 13 if max(len(code), len(out)) <= 42 else 11.5
        rect(slide, 0.92, 5.68, 5.45, 1.02, fill=PANEL)
        text(slide, 1.14, 5.82, 5.05, 0.36,
             [[('>>> ', dict(color=ACCENT, bold=True)), (code, dict(font=MONO))]],
             size=run_size)
        text(slide, 1.14, 6.24, 5.05, 0.36,
             [[(out, dict(font=MONO, color=ACCENT, bold=True))]], size=run_size)
    if page['math']:
        rect(slide, 7.0, 2.1, 5.42, 2.4, fill=PANEL, line_color=LINE, line_w=0.75)
        place_formula(slide, page['math'], cache_dir, 9.71, 3.3, 4.85, 1.85)
        spine_top = 4.85
    else:
        spine_top = 2.3
    term_spine(slide, page['visual'], 7.32, spine_top, 6.72)
    return slide


def image_slide(prs, lesson, page, image_path, cache_dir):
    slide = blank(prs)
    header(slide, lesson, page['title'])
    box_x, box_y, box_w, box_h = 0.92, 2.05, 7.55, 4.8
    with Image.open(image_path) as im:
        iw, ih = im.size
    scale = min(box_w / iw, box_h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = box_x + (box_w - pw) / 2, box_y + (box_h - ph) / 2
    slide.shapes.add_picture(str(image_path), Inches(px), Inches(py),
                             width=Inches(pw), height=Inches(ph))
    rect(slide, px, py, pw, ph, fill=None, line_color=LINE, line_w=1.0)
    text(slide, 8.85, 2.12, 3.55, 0.3, ['怎么读这张图'], size=11.5,
         color=ACCENT, bold=True)
    lines = [[(pt, {})] for pt in page['points'][:3]]
    text(slide, 8.85, 2.62, 3.55, 3.4, lines, size=15, line=1.32, space_after=11)
    nb, cell, _ = page['image']
    short = cell if len(cell) <= 12 else cell[:8] + '...'
    text(slide, 8.85, 6.42, 3.55, 0.3,
         [[(f'{nb}', dict(font=MONO)), (f'  ·  cell {short}', dict(font=MONO))]],
         size=10.5, color=SUB)
    return slide


def code_slide(prs, lesson):
    slide = blank(prs)
    header(slide, lesson, '代码实操')
    target = lesson['notebook'] or 'run_config.json / 精读笔记'
    text(slide, 0.92, 2.02, 8.4, 0.32,
         [[('打开  ', dict(color=SUB)), (target, dict(font=MONO, color=INK))]],
         size=13)
    lines = lesson['code'].split('\n')
    panel_h = min(0.72 + len(lines) * 0.345, 3.6)
    longest = max(len(line) for line in lines)
    code_size = 15 if longest <= 50 else (13 if longest <= 62 else 12)
    rect(slide, 0.92, 2.72, 7.5, panel_h, fill=CODE_BG,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.045)
    box, tf = _tf(slide, 1.34, 2.72 + 0.34, 6.75, panel_h - 0.62)
    for i, line in enumerate(lines):
        code_part, _, comment = line.partition('#')
        runs = [(code_part, dict(font=MONO, color=CODE_FG))]
        if comment:
            runs.append(('#' + comment, dict(font=MONO, color=CODE_COMMENT)))
        _para(tf, runs, first=(i == 0), line=1.42, size=code_size)
    text(slide, 0.92, 2.72 + panel_h + 0.28, 7.5, 0.35,
         [[('定位 cell  ', dict(color=SUB)),
           (', '.join(lesson['cells']), dict(font=MONO, color=INK))]], size=13)
    if lesson['notebook'] in {'00_environment_check.ipynb',
                              '15_imbalance_analysis_teaching.ipynb',
                              '16_protonet_teaching.ipynb',
                              '17_openset_teaching.ipynb'}:
        rail = ('演示入口', ['演示资产：results/teaching_ready/',
                            '载入后只做推理与检查', '训练与校准已在课前完成'])
    elif lesson['notebook']:
        rail = ('现场安排', ['现场：逐个运行本页 cell', '完整训练：课前排队完成',
                            '结果解读：以当前输出为准'])
    else:
        rail = ('研究入口', ['入口：run_config.json / metrics.json',
                            '对照 source-audit.md 逐条核验', '产出：受限结果段落'])
    text(slide, 8.85, 2.12, 3.55, 0.3, [rail[0]], size=11.5, color=ACCENT,
         bold=True)
    rail_lines = [[('—  ', dict(color=ACCENT, bold=True)), (pt, {})] for pt in rail[1]]
    text(slide, 8.85, 2.62, 3.55, 2.6, rail_lines, size=15, line=1.25,
         space_after=11)
    return slide


def exercise_slide(prs, lesson):
    slide = blank(prs)
    header(slide, lesson, '课堂练习')
    text(slide, 0.92, 2.35, 6.6, 2.6, [lesson['change']], size=23, line=1.35)
    rect(slide, 0.92, 5.35, 6.2, 0.66, fill=None, line_color=ACCENT,
         line_w=1.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(slide, 1.3, 5.53, 5.6, 0.35,
         [[(f"任务 {lesson['assignment']}", dict(bold=True, color=ACCENT)),
           ('　交付：代码 + 配置 + 解释', dict(color=INK))]], size=14)
    box_x, box_y, box_w, box_h = 8.15, 2.25, 4.25, 3.76
    rect(slide, box_x, box_y, box_w, box_h, fill=None, line_color=LINE,
         line_w=1.2)
    text(slide, box_x + 0.35, box_y + 0.32, box_w - 0.7, 0.35,
         ['先写预测，再运行'], size=13, bold=True, color=ACCENT)
    for i in range(3):
        hairline(slide, box_x + 0.35, box_y + 1.42 + i * 0.78, box_w - 0.7)
    return slide


def ai_slide(prs, lesson):
    slide = blank(prs)
    header(slide, lesson, 'AI 辅助')
    rect(slide, 0.92, 2.9, 7.55, 2.9, fill=PANEL)
    rect(slide, 0.92, 2.9, 0.055, 2.9, fill=ACCENT)
    text(slide, 1.35, 3.17, 6.85, 0.3, ['提示词（可直接复制）'], size=11.5,
         color=ACCENT, bold=True)
    text(slide, 1.35, 3.63, 6.85, 2.0, [lesson['ai']], size=15.5, line=1.4)
    text(slide, 8.85, 2.97, 3.55, 0.3, ['人工核验'], size=11.5, color=ACCENT,
         bold=True)
    text(slide, 8.85, 3.43, 3.55, 2.4, [lesson['verify']], size=14.5, line=1.35)
    return slide


def limits_slide(prs, lesson):
    slide = blank(prs)
    header(slide, lesson, '小结与边界')
    lines = [[('—  ', dict(color=ACCENT, bold=True)), (v, {})] for v in lesson['limits']]
    box = text(slide, 0.92, 2.35, 6.2, 3.6, lines, size=20, line=1.25,
               space_after=14)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    rect(slide, 7.9, 2.3, 4.5, 3.0, fill=PANEL)
    rect(slide, 7.9, 2.3, 0.055, 3.0, fill=ACCENT)
    text(slide, 8.32, 2.62, 3.85, 0.3, ['回到本课开头的问题'], size=11.5,
         color=ACCENT, bold=True)
    text(slide, 8.32, 3.1, 3.85, 2.0, [lesson['title']], size=19, line=1.4)
    return slide


def sources_slide(prs, lesson, sources):
    slide = blank(prs)
    header(slide, lesson, '来源、版本与复现')
    y = 2.18
    for key in lesson['sources']:
        name, url = sources[key]
        text(slide, 0.92, y, 4.4, 0.35, [name], size=15, bold=True)
        text(slide, 5.4, y + 0.03, 7.0, 0.32, [url], size=12, color=SUB)
        y += 0.72
        hairline(slide, 0.92, y - 0.18, 11.49)
    text(slide, 0.92, 6.3, 11.5, 0.4,
         ['完整来源状态与许可边界见 docs/teaching/source-audit.md；图像取自当前 Notebook 输出。'],
         size=12.5, color=SUB)
    return slide

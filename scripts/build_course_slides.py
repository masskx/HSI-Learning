"""Build all L00-L10 editable PPTX decks with the course design system.

Layout lives in slides/src/deck_design.py (python-pptx); content lives in
slides/src/course_content.py. Formulas are typeset via matplotlib mathtext
and embedded as transparent high-DPI images; LaTeX source is preserved in
course_content.py. The earlier ppt-generator-skill shell (v1 decks) was
replaced after visual review; the skill remains installed but unused.

No dependency installation, network upload, or rendering claims. Run with
the project Python, then verify with check_course_slides.py and export
previews with export_deck_previews.ps1.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
from pathlib import Path
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'slides/src'))
from course_content import LESSONS, SOURCES  # noqa: E402
import deck_design as D  # noqa: E402


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--only', nargs='+', choices=[v['id'] for v in LESSONS])
    args = parser.parse_args()
    try:
        from pptx import Presentation
        from PIL import Image
    except ImportError as exc:
        raise SystemExit(f'Missing PPT dependency: {exc}. Install requirements-slides.txt in the project environment.')

    slides_root = ROOT / 'slides'
    decks = slides_root / 'decks'
    assets = slides_root / 'assets/generated'
    formulas = slides_root / 'assets/formulas'
    decks.mkdir(parents=True, exist_ok=True)
    assets.mkdir(parents=True, exist_ok=True)
    source_records = {}
    built = []

    def image_asset(reference):
        notebook, cell_id, ordinal = reference
        path = ROOT / 'notebooks' / notebook
        nb = json.loads(path.read_text(encoding='utf-8'))
        if cell_id.startswith('cell') and cell_id[4:].isdigit():
            index = int(cell_id[4:])
            cells = [nb['cells'][index]] if 0 <= index < len(nb['cells']) else []
            cell_ref = f'position-{index} (legacy cell without id)'
        else:
            cells = [c for c in nb['cells'] if c.get('id') == cell_id]
            cell_ref = f'id-{cell_id}'
        if len(cells) != 1:
            raise ValueError(f'Missing/ambiguous image cell {notebook}:{cell_id}')
        outputs = [o['data']['image/png'] for o in cells[0].get('outputs', [])
                   if o.get('data', {}).get('image/png')]
        if ordinal >= len(outputs):
            raise ValueError(f'Missing embedded image {reference}; execute verified notebook first')
        raw = outputs[ordinal]
        raw = ''.join(raw) if isinstance(raw, list) else raw
        blob = base64.b64decode(raw)
        sha = hashlib.sha256(blob).hexdigest()
        filename = f'{Path(notebook).stem}-{cell_id}-{ordinal}-{sha[:10]}.png'
        target = assets / filename
        if target.exists() and target.read_bytes() != blob:
            raise ValueError('Image hash collision')
        target.write_bytes(blob)
        source_records.setdefault(filename, dict(
            kind='course-notebook-output', notebook=notebook, cell=cell_ref,
            sha256=sha, license='Course-generated figure; underlying dataset conditions still apply',
            path=str(target.relative_to(slides_root)), uses=[]))
        return target, filename

    for lesson in LESSONS:
        if args.only and lesson['id'] not in args.only:
            continue
        prs = Presentation()
        prs.slide_width = D.Inches(D.SLIDE_W)
        prs.slide_height = D.Inches(D.SLIDE_H)
        prs.core_properties.title = lesson['title']
        prs.core_properties.subject = 'Theory + practice + verified AI research assistance'
        prs.core_properties.author = 'HSI-Learning'
        records = []

        def notes(slide, content, source_ids=()):
            txt = content + '\n\n教学边界：本课不是榜单排名；历史结果不等同修正版实验。\n'
            txt += 'Notebook: ' + str(lesson['notebook']) + '\n定位: ' + ', '.join(lesson['cells']) + '\n'
            for sid in source_ids:
                txt += SOURCES[sid][0] + '\n' + SOURCES[sid][1] + '\n'
            slide.notes_slide.notes_text_frame.text = txt

        slide = D.cover(prs, lesson)
        notes(slide, '用真实问题开场，先让学生提出预测。不要从依赖列表开始读。')
        records.append({'type': 'cover', 'title': lesson['title']})

        slide = D.objectives(prs, lesson)
        notes(slide, '完成标准是能独立生成交付物，而不是只复述概念。右侧面板是本课的Notebook与cell入口。')
        records.append({'type': 'objectives', 'title': '本课交付'})

        for entry in lesson['pages']:
            index = len(prs.slides) + 1
            if entry['image']:
                path, key = image_asset(entry['image'])
                source_records[key]['uses'].append({'lesson': lesson['id'], 'slide': index})
                slide = D.image_slide(prs, lesson, entry, path, formulas)
            else:
                slide = D.theory(prs, lesson, entry, formulas)
            notes(slide, entry['notes'], [entry['source']] if entry['source'] else lesson['sources'][:1])
            records.append({'type': 'image' if entry['image'] else 'theory',
                            'title': entry['title'], 'image': entry['image'],
                            'math': entry['math']})

        slide = D.code_slide(prs, lesson)
        notes(slide, '这是关键代码节选，完整上下文在Notebook中。现场只做最小例或载入模型推理；完整训练提前排队。')
        records.append({'type': 'practice', 'title': '代码实操'})

        slide = D.exercise_slide(prs, lesson)
        notes(slide, '先让学生在右侧预测框写出预期，再运行一次受控改动。若与预期不同，保留记录而不是换种子直到符合故事。')
        records.append({'type': 'exercise', 'title': '课堂练习'})

        slide = D.ai_slide(prs, lesson)
        notes(slide, '可复制提示词见内容源。数据边界：不上传未授权数据或未公开稿件；不让AI虚构引用、指标或执行记录。'
               + lesson['ai'])
        records.append({'type': 'ai-assist', 'title': 'AI辅助'})

        slide = D.limits_slide(prs, lesson)
        notes(slide, '结束时请学生用一句话回答本课开头的问题，要求带证据类型。')
        records.append({'type': 'limits', 'title': '小结与边界'})

        slide = D.sources_slide(prs, lesson, SOURCES)
        notes(slide, '每次发布应补源码版本与运行配置；网络资料可访问不等于可自由再分发。', lesson['sources'])
        records.append({'type': 'sources', 'title': '来源与复现'})

        target = decks / f"{lesson['id']}-{lesson['slug']}.pptx"
        temp = target.with_suffix('.building.pptx')
        prs.save(temp)
        with zipfile.ZipFile(temp) as z:
            if z.testzip():
                raise ValueError(f'Corrupt pptx: {temp}')
        reloaded = Presentation(temp)
        if len(reloaded.slides) != len(records):
            raise ValueError('Slide count mismatch')
        temp.replace(target)
        built.append(dict(id=lesson['id'], title=lesson['title'],
                          path=str(target.relative_to(slides_root)),
                          slides=len(records), pages=records,
                          notebook=lesson['notebook'], cells=lesson['cells'],
                          generated=True, rendered=False, visual_review='pending',
                          sha256=hashlib.sha256(target.read_bytes()).hexdigest()))
        print(f"Generated {target.name}: {len(records)} editable slides", flush=True)

    report = dict(
        design=dict(system='slides/src/deck_design.py', engine='python-pptx',
                    fonts=dict(cjk='Microsoft YaHei (latin+ea set per run)', code='Consolas'),
                    formulas='matplotlib mathtext -> transparent PNG, LaTeX source in course_content.py',
                    motif='hyperspectral gradient strip; no footers, no page numbers',
                    replaces='ppt-generator-skill v1 shell (withdrawn after visual review)'),
        decks=built,
        note='PPTX generated and package-checked; rendering/visual review is a separate step.')
    suffix = '-' + '-'.join(args.only) if args.only else ''
    (slides_root / f'lesson-manifest{suffix}.json').write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    (slides_root / 'assets' / f'sources{suffix}.json').write_text(
        json.dumps(list(source_records.values()), ensure_ascii=False, indent=2),
        encoding='utf-8')
    print('No rendering claimed. Verify with check_course_slides.py and export_deck_previews.ps1.')


if __name__ == '__main__':
    main()
